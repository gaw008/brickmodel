"""Read the published normalized TGA observations; retain all source values.

Requires openpyxl (the extraction environment is recorded in the report).
The interval F1 calculation is a diagnostic of a declared simple rate form,
not an identification of an intrinsic surface or whole-brick reaction rate.
"""
import argparse
import csv
import json
import math
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--parameters', type=Path, required=True)
    args = parser.parse_args()
    settings = json.loads(args.parameters.read_text())
    source = Path(settings['source_directory'])
    out = Path(settings['output_directory'])
    book = openpyxl.load_workbook(source / settings['workbook'], data_only=True,
                                 read_only=True)
    baseline = settings['normalized_calcined_mass']
    capacity = settings['approximate_source_capacity_g_co2_per_g_sorbent']
    curves = []
    observations = []
    for sheet_name, atmosphere in settings['sheets'].items():
        sheet = book[sheet_name]
        rows = list(sheet.iter_rows(min_row=settings['data_first_row'], values_only=True))
        for cycle, column in enumerate(settings['cycle_columns'], start=1):
            records = []
            for index, row in enumerate(rows):
                item = {
                    'atmosphere': atmosphere, 'cycle': cycle,
                    'source_workbook': settings['workbook'], 'source_sheet': sheet_name,
                    'source_row': settings['data_first_row'] + index,
                    'source_mass_column': column,
                    'time_s': row[settings['time_column'] - 1],
                    'temperature_c': row[settings['temperature_column'] - 1],
                    'normalized_sample_mass': row[column - 1],
                    'uptake_g_co2_per_g_sorbent': row[column - 1] - baseline,
                }
                records.append(item)
            observations.extend(records)
            by_time = {r['time_s']: r for r in records}
            q = lambda time: by_time[time]['uptake_g_co2_per_g_sorbent']
            rates = []
            for start, end in settings['isothermal_rate_windows_s']:
                rate = math.log((capacity - q(start)) / (capacity - q(end))) / (end - start)
                rates.append({'start_s': start, 'end_s': end,
                              'conditional_interval_f1_rate_per_s': rate})
            values = [r['conditional_interval_f1_rate_per_s'] for r in rates]
            end_item = by_time[settings['source_sorption_end_s']]
            curves.append({
                'atmosphere': atmosphere, 'cycle': cycle, 'records': len(records),
                'strictly_increasing_time': all(b['time_s'] > a['time_s']
                                               for a, b in zip(records, records[1:])),
                'source_time_step_range_s': [min(b['time_s'] - a['time_s'] for a,b in zip(records,records[1:])),
                                             max(b['time_s'] - a['time_s'] for a,b in zip(records,records[1:]))],
                'temperature_at_and_before_sorption_end_matches_nominal': all(
                    r['temperature_c'] == settings['nominal_sorption_temperature_c']
                    for r in records if r['time_s'] <= settings['source_sorption_end_s']),
                'first_above_nominal_temperature': next(r for r in records if
                    r['temperature_c'] > settings['nominal_sorption_temperature_c']),
                'negative_uptake_count_retained': sum(r['uptake_g_co2_per_g_sorbent'] < 0 for r in records),
                'minimum_mass_record': min(records, key=lambda r:r['normalized_sample_mass']),
                'maximum_mass_record': max(records, key=lambda r:r['normalized_sample_mass']),
                'sorption_end_record': end_item,
                'end_uptake_over_approximate_source_capacity': end_item['uptake_g_co2_per_g_sorbent'] / capacity,
                'reported_samples': [by_time[t] for t in settings['report_times_s']],
                'conditional_f1_interval_rates': rates,
                'conditional_f1_interval_rate_max_over_min': max(values) / min(values),
            })
    target = out / 'observations.csv'
    with target.open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(observations[0]))
        writer.writeheader()
        writer.writerows(observations)
    # Numeric readback is part of the observation conversion, not a software test.
    with target.open(newline='') as stream:
        restored = list(csv.DictReader(stream))
    numeric_fields = ['time_s','temperature_c','normalized_sample_mass','uptake_g_co2_per_g_sorbent']
    max_roundtrip_error = {key: max(abs(float(a[key]) - b[key]) for a,b in
                                  zip(restored,observations,strict=True)) for key in numeric_fields}
    cross_cycle = {}
    for atmosphere in settings['sheets'].values():
        selected = [c for c in curves if c['atmosphere'] == atmosphere]
        first = selected[0]['sorption_end_record']['uptake_g_co2_per_g_sorbent']
        last = selected[-1]['sorption_end_record']['uptake_g_co2_per_g_sorbent']
        cross_cycle[atmosphere] = {'first_cycle_end_uptake':first,'last_cycle_end_uptake':last,
                                  'last_minus_first_uptake':last-first,
                                  'fractional_change_from_first':(last-first)/first}
    long = settings['long_exposure']
    long_book = openpyxl.load_workbook(source / long['workbook'],data_only=True,read_only=True)
    rows = list(long_book[long['sheet']].iter_rows(min_row=settings['data_first_row'],values_only=True))
    dry_rows = [r for r in rows if r[long['time_column']-1] <= long['caption_dry_exposure_end_s']]
    temperatures = sorted({r[long['temperature_column']-1] for r in dry_rows})
    result = {
        'parameters':settings, 'openpyxl_version':openpyxl.__version__,
        'observation_class':'author-published processed normalized measurements, not raw balance signal',
        'curves':curves, 'total_observations':len(observations),
        'csv_numeric_readback_maximum_difference':max_roundtrip_error,
        'cross_cycle_end_uptake':cross_cycle,
        'long_exposure_conflict': {
            'rows':len(rows), 'dry_exposure_rows':len(dry_rows),
            'caption_temperature_c':long['caption_temperature_c'],
            'workbook_dry_exposure_temperatures_c':temperatures,
            'caption_matches_workbook': temperatures == [long['caption_temperature_c']],
            'temperature_choice_for_prediction':None,
            'resolution':'Unresolved primary paper caption versus author workbook; excluded from any thermal rate inference.',
        },
        'kinetic_parameters_fitted':False, 'intrinsic_kinetics_qualified':False,
        'material_qualified':False,'training_eligible':False,
    }
    (out / 'review.json').write_text(json.dumps(result,indent=2,allow_nan=False)+'\n')
    print(json.dumps({'observations':len(observations),'readback':max_roundtrip_error,
                      'cross_cycle':cross_cycle,'long_exposure':result['long_exposure_conflict']}))


if __name__ == '__main__':
    main()
